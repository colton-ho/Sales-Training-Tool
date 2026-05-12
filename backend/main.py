import os
import uuid
import json
import shutil
import tempfile
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, File, UploadFile, HTTPException, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from pptx import Presentation
from openai import OpenAI

app = FastAPI(title="Sales Training Tool")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)

# In-memory session store
sessions: dict[str, dict] = {}

def get_client() -> OpenAI:
    key = os.environ.get("OPENAI_API_KEY")
    if not key:
        raise HTTPException(
            status_code=503,
            detail="伺服器未設定 OPENAI_API_KEY，請聯絡管理員",
        )
    return OpenAI(api_key=key)


def extract_ppt_text(ppt_path: Path) -> list[dict]:
    """Extract text content from each slide of a PowerPoint file."""
    prs = Presentation(str(ppt_path))
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = para.text.strip()
                    if line:
                        texts.append(line)
        slides.append({"slide": i, "content": texts})
    return slides


def build_analysis_prompt(slides: list[dict], transcript: str) -> str:
    slides_text = ""
    for s in slides:
        content = "\n".join(s["content"]) if s["content"] else "(無內容)"
        slides_text += f"--- 第 {s['slide']} 張投影片 ---\n{content}\n\n"

    return f"""你是一位專業的銷售培訓教練。以下是學員的銷售簡報投影片內容，以及學員練習簡報時的語音逐字稿。

請根據以下五個維度進行評估，每項滿分 20 分（總分 100 分）：
1. 內容完整性（Content Coverage）：是否涵蓋了所有投影片的重點
2. 語言表達（Language Quality）：用詞是否清晰、專業
3. 說服力（Persuasiveness）：是否能有效吸引客戶
4. 結構條理（Structure）：簡報邏輯是否清晰
5. 自信度（Confidence）：表達是否流暢、有自信

---

【投影片內容】
{slides_text}

【學員語音逐字稿】
{transcript}

---

請以 JSON 格式回覆，格式如下（請務必只回覆 JSON，不要加任何其他文字）：
{{
  "total_score": <整數 0-100>,
  "dimensions": {{
    "content_coverage": {{"score": <0-20>, "feedback": "<具體說明>"}},
    "language_quality": {{"score": <0-20>, "feedback": "<具體說明>"}},
    "persuasiveness": {{"score": <0-20>, "feedback": "<具體說明>"}},
    "structure": {{"score": <0-20>, "feedback": "<具體說明>"}},
    "confidence": {{"score": <0-20>, "feedback": "<具體說明>"}}
  }},
  "strengths": ["<優點1>", "<優點2>", "<優點3>"],
  "improvements": ["<改善建議1>", "<改善建議2>", "<改善建議3>"],
  "overall_comment": "<整體評語（2-3 句）>"
}}"""


@app.post("/api/upload-ppt")
async def upload_ppt(file: UploadFile = File(...)):
    """Upload a PowerPoint file and extract its content."""
    if not file.filename.lower().endswith((".ppt", ".pptx")):
        raise HTTPException(status_code=400, detail="請上傳 .ppt 或 .pptx 格式的檔案")

    session_id = str(uuid.uuid4())
    session_dir = UPLOAD_DIR / session_id
    session_dir.mkdir(parents=True, exist_ok=True)

    ppt_path = session_dir / file.filename
    with open(ppt_path, "wb") as f:
        shutil.copyfileobj(file.file, f)

    try:
        slides = extract_ppt_text(ppt_path)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"無法解析投影片：{str(e)}")

    sessions[session_id] = {"slides": slides, "status": "uploaded"}

    return {
        "session_id": session_id,
        "slide_count": len(slides),
        "slides": slides,
    }


@app.post("/api/analyze")
async def analyze(
    session_id: str = Form(...),
    audio: UploadFile = File(...),
):
    """Transcribe voice recording and analyze the sales presentation."""
    if session_id not in sessions:
        raise HTTPException(status_code=404, detail="找不到該練習階段，請重新上傳投影片")

    session = sessions[session_id]
    slides = session["slides"]

    # Save audio to temp file
    suffix = Path(audio.filename).suffix if audio.filename else ".webm"
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        shutil.copyfileobj(audio.file, tmp)
        tmp_path = tmp.name

    try:
        # Transcribe with Whisper
        with open(tmp_path, "rb") as audio_file:
            transcription = get_client().audio.transcriptions.create(
                model="whisper-1",
                file=audio_file,
                language="zh",
            )
        transcript = transcription.text
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"語音轉文字失敗：{str(e)}")
    finally:
        os.unlink(tmp_path)

    if not transcript.strip():
        raise HTTPException(status_code=422, detail="未偵測到語音內容，請重新錄製")

    # Analyze with GPT-4
    prompt = build_analysis_prompt(slides, transcript)
    try:
        response = get_client().chat.completions.create(
            model="gpt-4o",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.3,
        )
        raw = response.choices[0].message.content.strip()
        # Strip markdown code fences if present
        if raw.startswith("```"):
            parts = raw.split("```")
            raw = parts[1] if len(parts) > 1 else raw
            if raw.startswith("json"):
                raw = raw[4:]
        analysis = json.loads(raw)
    except json.JSONDecodeError:
        raise HTTPException(status_code=502, detail="AI 回覆格式異常，請重試")
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"AI 分析失敗：{str(e)}")

    sessions[session_id]["transcript"] = transcript
    sessions[session_id]["analysis"] = analysis
    sessions[session_id]["status"] = "done"

    return {
        "session_id": session_id,
        "transcript": transcript,
        "analysis": analysis,
    }


@app.get("/api/session/{session_id}")
async def get_session(session_id: str):
    if session_id not in sessions:
        raise HTTPException(status_code=404, detail="找不到該練習階段")
    return sessions[session_id]


# Serve frontend
frontend_dir = Path(__file__).parent.parent / "frontend"
if frontend_dir.exists():
    app.mount("/", StaticFiles(directory=str(frontend_dir), html=True), name="frontend")
